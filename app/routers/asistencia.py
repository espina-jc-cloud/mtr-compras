"""
Asistencia — Fase 1: carga rápida de la planilla + cierre del día.

La carga es a día vencido: la planilla de hoy se pasa mañana. Por eso la
pantalla de carga abre en AYER y no en hoy, y por eso la pantalla principal
muestra el cierre de un día terminado en vez de un estado en vivo.

Todo el cálculo vive en app/asistencia_calc.py. Acá solo se leen formularios,
se guardan marcaciones y se arma lo que ve la pantalla.
"""
import json
import os
import re
import tempfile
import uuid
from calendar import monthrange
from datetime import date, datetime, timedelta

from fastapi import APIRouter, Request, Depends, File, Form, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from app.database import get_db
from app.permissions import can, require_perm
from app.templates import templates
from app.asistencia_import import (aplicar_dia, buscar_persona, clasificar_dia,
                                   indexar_personas, normalizar, parsear)
from app.asistencia_calc import (
    agregar_bloque, get_config, guardar_bloque, jornada_esperada, limpiar_dia,
    recalcular_jornada,
)
from app.models_asistencia import (
    AsistenciaAuditLog, AsistenciaBloque, AsistenciaFeriado, AsistenciaJornada,
    AsistenciaImportacion, AsistenciaJornadaTipo, AsistenciaMotivo,
    AsistenciaPersona,
    ESTADO_EXTRA_CSS, ESTADO_REGISTRO_CSS, PLANTAS, TURNOS_OPERATIVO,
    hhmm_a_min, min_a_hhmm, min_a_texto,
)

router = APIRouter(prefix="/asistencia")

_ver    = require_perm("asistencia.ver")
_cargar = require_perm("asistencia.cargar")


# ══════════════════════════════════════════════════════════════════════════════
# Helpers
# ══════════════════════════════════════════════════════════════════════════════

def _parse_fecha(valor, default: date) -> date:
    if not valor:
        return default
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(valor).strip(), fmt).date()
        except ValueError:
            continue
    return default


def _fecha_carga_default(db: Session) -> date:
    """Ayer, salvo que la config diga otra cosa. Abrir en 'hoy' sería el error
    más probable del sistema: cargar la planilla de ayer sobre la fecha de hoy."""
    cfg = get_config(db)
    return date.today() - timedelta(days=cfg.dias_offset_carga or 0)


def _personas_activas(db: Session, planta: str = ""):
    """Nómina activa. Cada planta tiene su propia planilla de portería, así que
    la pantalla de carga se filtra por planta."""
    q = (db.query(AsistenciaPersona)
         .options(joinedload(AsistenciaPersona.proveedor))
         .filter(AsistenciaPersona.activo == True))  # noqa: E712
    if planta:
        q = q.filter(AsistenciaPersona.planta == planta)
    return q.order_by(AsistenciaPersona.orden_planilla,
                      AsistenciaPersona.apellido).all()


def _jornadas_por_persona(db: Session, fecha: date) -> dict:
    return {
        j.persona_id: j
        for j in db.query(AsistenciaJornada).filter(AsistenciaJornada.fecha == fecha).all()
    }


def _bloques_por_persona(db: Session, fecha: date) -> dict:
    out = {}
    for b in (db.query(AsistenciaBloque)
              .filter(AsistenciaBloque.fecha_operativa == fecha)
              .order_by(AsistenciaBloque.desde).all()):
        out.setdefault(b.persona_id, []).append(b)
    return out


def _fila_ctx(db, persona, fecha, jornada=None, bloques=None):
    """Contexto de una fila de la pantalla de carga."""
    esp = jornada_esperada(db, persona, fecha)
    if bloques is None:
        bloques = (
            db.query(AsistenciaBloque)
            .filter(AsistenciaBloque.persona_id == persona.id,
                    AsistenciaBloque.fecha_operativa == fecha)
            .order_by(AsistenciaBloque.desde).all()
        )
    return {
        "p": persona, "fecha": fecha, "j": jornada, "bloques": bloques, "esp": esp,
    }


def _estado_mes(db: Session, anio: int, mes: int, personas_esperadas: int) -> list:
    """Un renglón por día del mes con su estado de carga.

    Es la pieza que hace usable la carga de un mes entero: sin esto, el riesgo
    real de la carga diferida no es no ver el presente, es que se pase un día
    sin cargar y el total del mes quede mal.
    """
    total_dias = monthrange(anio, mes)[1]
    hoy = date.today()

    cargados = dict(
        db.query(AsistenciaJornada.fecha, func.count(AsistenciaJornada.id))
        .filter(AsistenciaJornada.fecha >= date(anio, mes, 1),
                AsistenciaJornada.fecha <= date(anio, mes, total_dias))
        .group_by(AsistenciaJornada.fecha).all()
    )
    feriados = {
        f.fecha: f.nombre
        for f in db.query(AsistenciaFeriado).filter(
            AsistenciaFeriado.fecha >= date(anio, mes, 1),
            AsistenciaFeriado.fecha <= date(anio, mes, total_dias)).all()
    }

    dias = []
    for d in range(1, total_dias + 1):
        f = date(anio, mes, d)
        n = cargados.get(f, 0)
        if f > hoy:
            estado = "futuro"
        elif f.weekday() == 6 or f in feriados:
            estado = "cargado" if n else "no_laborable"
        elif n == 0:
            estado = "vacio"
        elif n < personas_esperadas:
            estado = "incompleto"
        else:
            estado = "cargado"
        dias.append({
            "fecha": f, "dia": d, "n": n, "estado": estado,
            "feriado": feriados.get(f),
            "dow": f.weekday(),
        })
    return dias


# ══════════════════════════════════════════════════════════════════════════════
# Pantalla principal — el cierre del día
# ══════════════════════════════════════════════════════════════════════════════

@router.get("", response_class=HTMLResponse)
async def dia(request: Request, db: Session = Depends(get_db),
              current_user=Depends(_ver), fecha: str = "", planta: str = ""):
    f = _parse_fecha(fecha, _fecha_carga_default(db))

    q = (db.query(AsistenciaJornada)
         .options(joinedload(AsistenciaJornada.persona),
                  joinedload(AsistenciaJornada.motivo))
         .filter(AsistenciaJornada.fecha == f))
    if planta:
        q = q.filter(AsistenciaJornada.planta_snap == planta)
    jornadas = q.all()
    bloques = _bloques_por_persona(db, f)
    for j in jornadas:
        j._bloques = bloques.get(j.persona_id, [])

    mtr     = [j for j in jornadas if j.tipo_persona_snap != "tercero"]
    terceros = [j for j in jornadas if j.tipo_persona_snap == "tercero"]
    trabajaron = [j for j in mtr if j.cantidad_bloques > 0]

    # Acumulado del mes, hasta la fecha vista inclusive.
    ini_mes = f.replace(day=1)
    extra_mes = (
        db.query(func.coalesce(func.sum(AsistenciaJornada.minutos_extra_detectada), 0))
        .filter(AsistenciaJornada.fecha >= ini_mes, AsistenciaJornada.fecha <= f,
                AsistenciaJornada.tipo_persona_snap != "tercero").scalar() or 0
    )

    kpis = {
        "trabajaron":   len(trabajaron),
        "ausentes":     len([j for j in mtr if j.estado_registro == "ausente"]),
        "extra":        sum(j.minutos_extra_detectada or 0 for j in mtr),
        "e50":          sum(j.minutos_extra_50 or 0 for j in mtr),
        "e100":         sum(j.minutos_extra_100 or 0 for j in mtr),
        "fuera":        sum(j.minutos_fuera_horario or 0 for j in mtr),
        "defecto":      sum(j.minutos_defecto or 0 for j in mtr),
        "extra_mes":    extra_mes,
        "terceros":     len(terceros),
        "horas_terceros": sum(j.minutos_presencia or 0 for j in terceros),
        "personas_con_extra": len([j for j in mtr if (j.minutos_extra_detectada or 0) > 0]),
    }

    # ── Requiere atención ────────────────────────────────────────────────────
    # Cada ítem contesta tres cosas: qué pasa, por qué importa, y dónde se
    # resuelve. Un aviso que no lleva a la acción no sirve de nada.
    qs_planta = f"&planta={planta}" if planta else ""
    mes_qs = f.strftime("%Y-%m")

    def _acumulado_mes(persona_id):
        return (db.query(func.coalesce(
            func.sum(AsistenciaJornada.minutos_extra_detectada), 0))
            .filter(AsistenciaJornada.persona_id == persona_id,
                    AsistenciaJornada.fecha >= ini_mes,
                    AsistenciaJornada.fecha <= f).scalar() or 0)

    atencion = []
    for j in jornadas:
        nom = f"{j.persona.apellido}, {j.persona.nombre}" if j.persona else "?"
        url_carga = (f"/asistencia/carga?fecha={f.isoformat()}{qs_planta}"
                     f"&foco={j.persona_id}#fila-{j.persona_id}")
        url_ficha = f"/asistencia/persona/{j.persona_id}?mes={mes_qs}&foco={f.isoformat()}"

        if j.estado_registro == "inconsistente":
            atencion.append((0, j, "Registro inconsistente",
                             "Hay un egreso sin ingreso, dos ingresos seguidos o un bloque "
                             "de más de 16 h. Mientras quede así, las horas de este día no "
                             "computan ni como extra ni como jornada cumplida.",
                             url_carga, "Corregir la fila"))
        elif j.estado_registro == "sin_egreso":
            atencion.append((1, j, "Falta el egreso",
                             f"{nom} tiene la entrada cargada pero no la salida. No se le "
                             "computa extra ni horas no cumplidas: el sistema no puede "
                             "inventar a qué hora se fue.",
                             url_carga, "Cargar el egreso"))
        elif j.revisar:
            atencion.append((2, j, "Cambió después de resolverse",
                             "El recálculo dio distinto de lo que había cuando alguien "
                             "resolvió esta extra. La decisión NO se tocó: hay que mirarla "
                             "de nuevo y confirmarla o corregirla.",
                             url_ficha, "Revisar la ficha"))
        elif (j.minutos_extra_detectada or 0) > 0 and j.estado_extra == "pendiente":
            acum = _acumulado_mes(j.persona_id)
            porque = (f"{min_a_texto(j.minutos_extra_detectada)} sin resolver. "
                      f"Lleva {min_a_texto(acum)} en lo que va del mes")
            if j.minutos_extra_100:
                porque += f", {min_a_texto(j.minutos_extra_100)} de ellas al 100 %"
            porque += ". Falta decidir el motivo y si se aprueba."
            atencion.append((3, j, "Extra sin resolver", porque,
                             url_ficha, "Analizar a la persona"))

    atencion.sort(key=lambda t: (t[0], -(t[1].minutos_extra_detectada or 0)))

    qp = db.query(func.count(AsistenciaPersona.id)).filter(
        AsistenciaPersona.activo == True,  # noqa: E712
        AsistenciaPersona.tipo == "mtr")
    if planta:
        qp = qp.filter(AsistenciaPersona.planta == planta)
    personas_activas = qp.scalar() or 0

    return templates.TemplateResponse(request, "asistencia/dia.html", {
        "current_user": current_user,
        "fecha": f,
        "planta": planta,
        "plantas": PLANTAS,
        "prev": f - timedelta(days=1),
        "next": f + timedelta(days=1),
        "hoy": date.today(),
        "jornadas": sorted(jornadas, key=lambda j: (
            -(j.minutos_extra_detectada or 0), j.persona.apellido if j.persona else "")),
        "atencion": atencion,
        "kpis": kpis,
        "cargadas": len(jornadas),
        "personas_activas": personas_activas,
        "dias_mes": _estado_mes(db, f.year, f.month, personas_activas),
        "estado_extra_css": ESTADO_EXTRA_CSS,
        "estado_registro_css": ESTADO_REGISTRO_CSS,
        "mt": min_a_texto,
    })


# ══════════════════════════════════════════════════════════════════════════════
# Carga rápida
# ══════════════════════════════════════════════════════════════════════════════

@router.get("/carga", response_class=HTMLResponse)
async def carga(request: Request, db: Session = Depends(get_db),
                current_user=Depends(_cargar), fecha: str = "", planta: str = "",
                foco: int = 0):
    f = _parse_fecha(fecha, _fecha_carga_default(db))

    personas = _personas_activas(db, planta)
    jornadas = _jornadas_por_persona(db, f)
    bloques  = _bloques_por_persona(db, f)

    filas = [_fila_ctx(db, p, f, jornadas.get(p.id), bloques.get(p.id, [])) for p in personas]

    cargadas = len([1 for x in filas if x["j"] is not None])
    sin_egreso = len([1 for x in filas if x["j"] is not None
                      and x["j"].estado_registro == "sin_egreso"])

    personas_mtr = len([p for p in personas if p.tipo == "mtr"])

    return templates.TemplateResponse(request, "asistencia/carga.html", {
        "current_user": current_user,
        "fecha": f,
        "planta": planta,
        "plantas": PLANTAS,
        "prev": f - timedelta(days=1),
        "next": f + timedelta(days=1),
        "hoy": date.today(),
        "filas": filas,
        "foco": foco,
        "total": len(filas),
        "cargadas": cargadas,
        "sin_egreso": sin_egreso,
        "dias_mes": _estado_mes(db, f.year, f.month, personas_mtr),
        "turnos": TURNOS_OPERATIVO,
        "mt": min_a_texto,
    })


@router.post("/carga/fila", response_class=HTMLResponse)
async def guardar_fila(request: Request, db: Session = Depends(get_db),
                       current_user=Depends(_cargar),
                       persona_id: int = Form(...), fecha: str = Form(...),
                       ingreso: str = Form(""), egreso: str = Form("")):
    """Guarda una fila. Devuelve la fila re-renderizada (HTMX).

    Guardado por fila y no por formulario completo: si se corta la conexión o
    se cierra el navegador a mitad de una planilla, no se pierde nada.
    """
    f = _parse_fecha(fecha, date.today())
    persona = db.query(AsistenciaPersona).get(persona_id)
    if persona is None:
        return HTMLResponse("", status_code=404)

    ingreso, egreso = (ingreso or "").strip(), (egreso or "").strip()

    if not ingreso and not egreso:
        limpiar_dia(db, persona, f, current_user.id)
        db.commit()
        return _render_fila(request, db, persona, f, current_user)

    guardar_bloque(db, persona, f, ingreso, egreso, user_id=current_user.id)
    recalcular_jornada(db, persona, f)
    db.commit()
    return _render_fila(request, db, persona, f, current_user)


@router.post("/carga/bloque", response_class=HTMLResponse)
async def sumar_bloque(request: Request, db: Session = Depends(get_db),
                       current_user=Depends(_cargar),
                       persona_id: int = Form(...), fecha: str = Form(...),
                       ingreso: str = Form(""), egreso: str = Form(""),
                       turno: str = Form("")):
    """Agrega un bloque más al día sin pisar el anterior.

    Es el caso del operativo de buque: trabajó 08-16, se fue, volvió 22-06.
    """
    f = _parse_fecha(fecha, date.today())
    persona = db.query(AsistenciaPersona).get(persona_id)
    if persona is None or not (ingreso or "").strip():
        return _render_fila(request, db, persona, f, current_user) if persona else HTMLResponse("", 404)

    agregar_bloque(db, persona, f, ingreso.strip(), (egreso or "").strip(),
                   user_id=current_user.id,
                   fuente="turno_operativo" if turno else "planilla",
                   turno=turno or None)
    recalcular_jornada(db, persona, f)
    db.commit()
    return _render_fila(request, db, persona, f, current_user)


@router.post("/carga/ausente", response_class=HTMLResponse)
async def marcar_ausente(request: Request, db: Session = Depends(get_db),
                         current_user=Depends(_cargar),
                         persona_id: int = Form(...), fecha: str = Form(...)):
    """Ausente explícito. Vacío significa 'todavía no cargué', no 'no vino':
    sin esa distinción no se puede saber si la planilla está completa."""
    f = _parse_fecha(fecha, date.today())
    persona = db.query(AsistenciaPersona).get(persona_id)
    if persona is None:
        return HTMLResponse("", status_code=404)
    limpiar_dia(db, persona, f, current_user.id)
    recalcular_jornada(db, persona, f, marcar_ausente=True)
    db.commit()
    return _render_fila(request, db, persona, f, current_user)


@router.post("/carga/masivo")
async def carga_masiva(db: Session = Depends(get_db), current_user=Depends(_cargar),
                       fecha: str = Form(...), accion: str = Form(...),
                       hora: str = Form(""), planta: str = Form("")):
    """Atajos de masa. En la planilla real 15 de 18 personas entran 08:00:
    este botón ahorra más tiempo que cualquier otra optimización."""
    f = _parse_fecha(fecha, date.today())
    personas = [p for p in _personas_activas(db, planta) if p.tipo == "mtr"]
    jornadas = _jornadas_por_persona(db, f)
    bloques  = _bloques_por_persona(db, f)
    n = 0

    for p in personas:
        j = jornadas.get(p.id)
        bs = bloques.get(p.id, [])
        esp = jornada_esperada(db, p, f)

        if accion == "ingresos":
            if bs or (j and j.estado_registro == "ausente"):
                continue
            h = hora or esp["hora_entrada"] or "08:00"
            guardar_bloque(db, p, f, h, "", user_id=current_user.id)
        elif accion == "egresos":
            if not bs or bs[-1].hasta is not None:
                continue
            h = hora or esp["hora_salida"] or "16:00"
            ing = bs[-1].desde.strftime("%H:%M")
            guardar_bloque(db, p, f, ing, h, user_id=current_user.id)
        elif accion == "ausentes":
            if bs or j is not None:
                continue
            limpiar_dia(db, p, f, current_user.id)
            recalcular_jornada(db, p, f, marcar_ausente=True)
            n += 1
            continue
        else:
            continue

        recalcular_jornada(db, p, f)
        n += 1

    db.commit()
    return RedirectResponse(
        url=f"/asistencia/carga?fecha={f.isoformat()}"
            f"{'&planta=' + planta if planta else ''}&ok={n}+filas+actualizadas",
        status_code=303)


def _render_fila(request, db, persona, fecha, current_user):
    j = (db.query(AsistenciaJornada)
         .filter(AsistenciaJornada.persona_id == persona.id,
                 AsistenciaJornada.fecha == fecha).first())
    ctx = _fila_ctx(db, persona, fecha, j)
    ctx.update({"current_user": current_user, "mt": min_a_texto,
                "turnos": TURNOS_OPERATIVO})
    return templates.TemplateResponse(request, "asistencia/_fila.html", ctx)


# ══════════════════════════════════════════════════════════════════════════════
# Nómina — administración del personal propio
# ══════════════════════════════════════════════════════════════════════════════

_config = require_perm("asistencia.configurar")


def _render_persona(request, db, persona, current_user):
    return templates.TemplateResponse(request, "asistencia/_persona_fila.html", {
        "current_user": current_user,
        "p": persona,
        "plantas": PLANTAS,
        "jornadas": db.query(AsistenciaJornadaTipo).filter(
            AsistenciaJornadaTipo.activo == True).order_by(AsistenciaJornadaTipo.id).all(),  # noqa: E712
    })


@router.get("/personal", response_class=HTMLResponse)
async def personal(request: Request, db: Session = Depends(get_db),
                   current_user=Depends(_config), ver_bajas: int = 0):
    q = db.query(AsistenciaPersona)
    if not ver_bajas:
        q = q.filter(AsistenciaPersona.activo == True)  # noqa: E712
    personas = q.order_by(AsistenciaPersona.orden_planilla,
                          AsistenciaPersona.apellido).all()

    jornadas = (db.query(AsistenciaJornadaTipo)
                .filter(AsistenciaJornadaTipo.activo == True)  # noqa: E712
                .order_by(AsistenciaJornadaTipo.id).all())

    propias = [p for p in personas if p.tipo == "mtr"]
    activos = [p for p in propias if p.activo]
    return templates.TemplateResponse(request, "asistencia/personal.html", {
        "current_user": current_user,
        "personas": propias,
        "plantas": PLANTAS,
        "jornadas": jornadas,
        "ver_bajas": ver_bajas,
        "sin_planta": len([p for p in activos if not p.planta]),
        "activos": len(activos),
        "uso_planta": {cod: len([p for p in activos if p.planta == cod])
                       for cod, _ in PLANTAS},
    })


@router.post("/personal/fila", response_class=HTMLResponse)
async def personal_guardar(request: Request, db: Session = Depends(get_db),
                           current_user=Depends(_config),
                           persona_id: int = Form(...),
                           planta: str = Form(""), jornada_tipo_id: str = Form(""),
                           legajo: str = Form(""), dni: str = Form(""),
                           orden_planilla: str = Form("")):
    """Guarda una fila de la nómina. Cada cambio queda auditado con su valor
    anterior: cambiar la planta o la jornada de alguien mueve números de meses
    enteros, así que no puede pasar sin registro."""
    p = db.query(AsistenciaPersona).get(persona_id)
    if p is None:
        return HTMLResponse("", status_code=404)

    cambios = [
        ("planta",          p.planta,          (planta or "").strip() or None),
        ("jornada_tipo_id", p.jornada_tipo_id, int(jornada_tipo_id) if jornada_tipo_id else None),
        ("legajo",          p.legajo,          (legajo or "").strip() or None),
        ("dni",             p.dni,             (dni or "").strip() or None),
        ("orden_planilla",  p.orden_planilla,  int(orden_planilla) if (orden_planilla or "").strip().isdigit() else p.orden_planilla),
    ]
    for campo, viejo, nuevo in cambios:
        if viejo == nuevo:
            continue
        setattr(p, campo, nuevo)
        db.add(AsistenciaAuditLog(
            entidad="persona", entidad_id=p.id, accion="editar", campo=campo,
            valor_anterior=str(viejo) if viejo is not None else None,
            valor_nuevo=str(nuevo) if nuevo is not None else None,
            user_id=current_user.id,
        ))
    db.commit()
    db.refresh(p)
    return _render_persona(request, db, p, current_user)


@router.post("/personal/estado", response_class=HTMLResponse)
async def personal_estado(request: Request, db: Session = Depends(get_db),
                          current_user=Depends(_config),
                          persona_id: int = Form(...)):
    """Alta / baja. NUNCA se borra una persona: sale de la planilla de carga
    pero queda en todos los históricos y reportes."""
    p = db.query(AsistenciaPersona).get(persona_id)
    if p is None:
        return HTMLResponse("", status_code=404)
    p.activo = not p.activo
    p.fecha_baja = None if p.activo else date.today()
    db.add(AsistenciaAuditLog(
        entidad="persona", entidad_id=p.id,
        accion="alta" if p.activo else "baja", campo="activo",
        valor_anterior=str(not p.activo), valor_nuevo=str(p.activo),
        user_id=current_user.id,
    ))
    db.commit()
    db.refresh(p)
    return _render_persona(request, db, p, current_user)


@router.post("/personal/nueva")
async def personal_nueva(db: Session = Depends(get_db), current_user=Depends(_config),
                         apellido: str = Form(...), nombre: str = Form(...),
                         planta: str = Form(""), jornada_tipo_id: str = Form(""),
                         legajo: str = Form(""), dni: str = Form("")):
    apellido, nombre = apellido.strip(), nombre.strip()
    if not apellido or not nombre:
        return RedirectResponse("/asistencia/personal?err=Falta+apellido+o+nombre", 303)

    ya = (db.query(AsistenciaPersona)
          .filter(AsistenciaPersona.apellido == apellido,
                  AsistenciaPersona.nombre == nombre).first())
    if ya is not None:
        if not ya.activo:   # reingreso: se reactiva, no se duplica
            ya.activo, ya.fecha_baja = True, None
            db.commit()
            return RedirectResponse("/asistencia/personal?ok=Persona+reactivada", 303)
        return RedirectResponse("/asistencia/personal?err=Esa+persona+ya+está+en+la+nómina", 303)

    ultimo = (db.query(func.coalesce(func.max(AsistenciaPersona.orden_planilla), 0))
              .scalar() or 0)
    p = AsistenciaPersona(
        apellido=apellido, nombre=nombre, tipo="mtr",
        planta=(planta or "").strip() or None,
        jornada_tipo_id=int(jornada_tipo_id) if jornada_tipo_id else None,
        legajo=(legajo or "").strip() or None, dni=(dni or "").strip() or None,
        orden_planilla=ultimo + 10, activo=True, fecha_alta=date.today(),
    )
    db.add(p)
    db.flush()
    db.add(AsistenciaAuditLog(entidad="persona", entidad_id=p.id, accion="crear",
                              valor_nuevo=f"{apellido}, {nombre}", user_id=current_user.id))
    db.commit()
    return RedirectResponse(f"/asistencia/personal?ok={apellido}+agregado+a+la+nómina", 303)


# ══════════════════════════════════════════════════════════════════════════════
# Importación del Excel diario de horas
# ══════════════════════════════════════════════════════════════════════════════

_IMPORT_DIR = os.path.join(tempfile.gettempdir(), "mtr_asistencia_import")


def _planta_de_hoja(hoja: str) -> str:
    """'HS. SEPTIEMBRE. MTR I 2026' → 'MTR1'. Devuelve '' si no se puede inferir.

    El nombre de la hoja dice de qué planta es la planilla, y esa es la planta
    que vale para la jornada: Báez rota entre plantas y en la planilla de MTR1
    sus horas son de MTR1, no las del maestro.
    """
    h = (hoja or "").upper().replace(".", " ")
    if re.search(r"MTR\s*II\b|MTR\s*2\b", h):
        return "MTR2"
    if re.search(r"MTR\s*I\b|MTR\s*1\b", h):
        return "MTR1"
    return ""


def _guardar_temporal(contenido: bytes) -> str:
    """Guarda el archivo subido y devuelve un token.

    El preview y la confirmación son dos requests distintos y el navegador no
    puede reenviar el archivo solo. Se guarda en el temporal del sistema: si se
    perdió (reinicio del server, filesystem efímero de Railway), se le pide al
    usuario que lo suba de nuevo en vez de fallar con un error críptico.
    """
    os.makedirs(_IMPORT_DIR, exist_ok=True)
    token = uuid.uuid4().hex
    with open(os.path.join(_IMPORT_DIR, token + ".xlsx"), "wb") as fh:
        fh.write(contenido)
    return token


def _leer_temporal(token: str):
    ruta = os.path.join(_IMPORT_DIR, os.path.basename(token) + ".xlsx")
    if not os.path.exists(ruta):
        return None
    with open(ruta, "rb") as fh:
        return fh.read()


def _analizar(db: Session, datos: dict, tipo: str = "mtr") -> dict:
    """Cruza lo parseado contra la nómina y contra lo ya cargado.

    El match se hace SOLO contra personas del mismo tipo: la planilla de
    terceros no puede resolver un nombre contra un empleado propio.
    """
    personas = db.query(AsistenciaPersona).filter(
        AsistenciaPersona.activo == True,  # noqa: E712
        AsistenciaPersona.tipo == tipo).all()
    idx = indexar_personas(personas)
    grupos = {}

    dias, no_reconocidos = [], {}
    for d in datos["dias"]:
        existentes = {
            j.persona_id: j for j in db.query(AsistenciaJornada)
            .filter(AsistenciaJornada.fecha == d["fecha"]).all()
        }
        bloques = _bloques_por_persona(db, d["fecha"])

        filas, iguales, distintas, nuevas = [], 0, 0, 0
        for f in d["filas"]:
            grupos[f["grupo"]] = grupos.get(f["grupo"], 0) + 1
            persona, motivo = buscar_persona(idx, f["apellido"], f["nombre"])
            if persona is None:
                clave = f"{f['apellido']}, {f['nombre']}"
                no_reconocidos.setdefault(clave, {"motivo": motivo, "veces": 0,
                                                  "grupo": f["grupo"]})
                no_reconocidos[clave]["veces"] += 1
                filas.append({**f, "persona": None, "estado": "sin_match"})
                continue

            bs = bloques.get(persona.id, [])
            actual_in = bs[0].desde.strftime("%H:%M") if bs else None
            actual_out = (bs[0].hasta.strftime("%H:%M")
                          if bs and bs[0].hasta else None)
            # El Excel escribe la medianoche como 24:00 y el sistema la guarda
            # como 00:00 del día siguiente: son la misma hora.
            nuevo_out = f["egreso"]
            if nuevo_out == "24:00" and actual_out == "00:00":
                nuevo_out = "00:00"

            if persona.id not in existentes:
                estado = "nueva"; nuevas += 1
            elif f["ausente"]:
                estado = ("igual" if existentes[persona.id].estado_registro == "ausente"
                          else "distinta")
                iguales += (estado == "igual"); distintas += (estado == "distinta")
            elif actual_in == f["ingreso"] and actual_out == nuevo_out:
                estado = "igual"; iguales += 1
            else:
                estado = "distinta"; distintas += 1

            filas.append({**f, "persona": persona, "estado": estado,
                          "actual_in": actual_in, "actual_out": actual_out})

        dias.append({
            "fecha": d["fecha"],
            "filas": filas,
            "n": len(filas),
            "cargadas": len(existentes),
            "iguales": iguales, "distintas": distintas, "nuevas": nuevas,
            "cambia": distintas + nuevas,
        })

    return {"dias": dias, "no_reconocidos": no_reconocidos, "grupos": grupos,
            "avisos": datos.get("avisos", [])}


@router.get("/importar", response_class=HTMLResponse)
async def importar_form(request: Request, db: Session = Depends(get_db),
                        current_user=Depends(_cargar), tipo: str = "mtr"):
    return templates.TemplateResponse(request, "asistencia/importar.html", {
        "current_user": current_user, "paso": "subir", "tipo_sel": tipo,
        "buzon": _estado_buzon(db),
    })


@router.post("/importar", response_class=HTMLResponse)
async def importar_preview(request: Request, db: Session = Depends(get_db),
                           current_user=Depends(_cargar),
                           archivo: UploadFile = File(None),
                           token: str = Form(""), hoja: str = Form(""),
                           tipo: str = Form("mtr"), planta: str = Form("")):
    """Vista previa. NO toca la base: solo muestra qué cambiaría."""
    if archivo is not None and archivo.filename:
        contenido = await archivo.read()
        token = _guardar_temporal(contenido)
    else:
        contenido = _leer_temporal(token) if token else None

    if not contenido:
        return templates.TemplateResponse(request, "asistencia/importar.html", {
            "current_user": current_user, "paso": "subir", "tipo_sel": tipo,
            "error": "Se perdió el archivo subido. Volvé a elegirlo.",
        })

    try:
        datos = parsear(contenido, hoja or None)
    except Exception as e:
        return templates.TemplateResponse(request, "asistencia/importar.html", {
            "current_user": current_user, "paso": "subir", "tipo_sel": tipo,
            "error": f"No pude leer el Excel: {e}",
        })

    analisis = _analizar(db, datos, tipo)
    return templates.TemplateResponse(request, "asistencia/importar.html", {
        "current_user": current_user, "paso": "previo",
        "token": token, "hoja": datos["hoja"], "hojas": datos["hojas"],
        "tipo": tipo, "plantas": PLANTAS,
        "planta": planta or _planta_de_hoja(datos["hoja"]),
        "nombre_archivo": archivo.filename if archivo and archivo.filename else "",
        **analisis,
    })


@router.post("/importar/aplicar")
async def importar_aplicar(db: Session = Depends(get_db), current_user=Depends(_cargar),
                           token: str = Form(...), hoja: str = Form(""),
                           fechas: list = Form([]), tipo: str = Form("mtr"),
                           planta: str = Form(""), alta_nuevos: str = Form("")):
    """Aplica los días tildados. Idempotente: reimportar un día reemplaza sus
    marcaciones (las viejas quedan anuladas, no borradas)."""
    contenido = _leer_temporal(token)
    if not contenido:
        return RedirectResponse(
            "/asistencia/importar?err=Se+perdió+el+archivo.+Subilo+de+nuevo", 303)

    datos = parsear(contenido, hoja or None)
    elegidas = {_parse_fecha(f, None) for f in fechas}
    elegidas.discard(None)
    if not elegidas:
        return RedirectResponse("/asistencia/importar?err=No+elegiste+ningún+día", 303)

    planta = planta or _planta_de_hoja(datos["hoja"])
    personas = db.query(AsistenciaPersona).filter(
        AsistenciaPersona.activo == True,  # noqa: E712
        AsistenciaPersona.tipo == tipo).all()
    idx = indexar_personas(personas)

    n_filas, n_dias, sin_match, n_altas = 0, 0, 0, 0
    antes = db.query(func.count(AsistenciaPersona.id)).scalar() or 0
    for d in datos["dias"]:
        if d["fecha"] not in elegidas:
            continue
        n_dias += 1
        a, sm = aplicar_dia(db, d["fecha"], d["filas"], idx, planta=planta,
                            tipo=tipo, user_id=current_user.id,
                            alta_nuevos=bool(alta_nuevos))
        n_filas += a
        sin_match += sm
        db.add(AsistenciaAuditLog(
            entidad="jornada", accion="importar", campo="dia",
            valor_nuevo=d["fecha"].isoformat(),
            motivo=f"Importación del Excel · hoja {datos['hoja']}",
            user_id=current_user.id,
        ))
    n_altas = (db.query(func.count(AsistenciaPersona.id)).scalar() or 0) - antes

    db.add(AsistenciaImportacion(
        origen="manual", estado="ok", asunto=None, hoja=datos["hoja"],
        dias_aplicados=n_dias, filas_aplicadas=n_filas,
        sin_reconocer=sin_match, usuario_id=current_user.id,
    ))
    db.commit()

    msg = f"{n_dias}+días+importados+·+{n_filas}+filas"
    if n_altas:
        msg += f"+·+{n_altas}+altas"
    if sin_match:
        msg += f"+·+{sin_match}+sin+reconocer"
    destino = "/asistencia/terceros" if tipo == "tercero" else "/asistencia"
    return RedirectResponse(f"{destino}?ok={msg}", 303)


@router.get("/terceros", response_class=HTMLResponse)
async def terceros(request: Request, db: Session = Depends(get_db),
                   current_user=Depends(_ver), mes: str = ""):
    """Personal de terceros — pantalla SEPARADA a propósito.

    Sus horas no se suman en ningún KPI junto a las del personal propio. Acá se
    ven por proveedor (el grupo de su planilla) y por persona.
    """
    hoy = date.today()
    try:
        anio, m = (int(x) for x in mes.split("-")) if mes else (hoy.year, hoy.month)
    except ValueError:
        anio, m = hoy.year, hoy.month
    desde, hasta = date(anio, m, 1), date(anio, m, monthrange(anio, m)[1])

    jornadas = (
        db.query(AsistenciaJornada)
        .options(joinedload(AsistenciaJornada.persona))
        .filter(AsistenciaJornada.fecha >= desde, AsistenciaJornada.fecha <= hasta,
                AsistenciaJornada.tipo_persona_snap == "tercero")
        .order_by(AsistenciaJornada.fecha).all()
    )

    por_grupo, por_persona = {}, {}
    for j in jornadas:
        g = j.grupo_snap or j.proveedor_snap or "Sin proveedor"
        gd = por_grupo.setdefault(g, {"minutos": 0, "dias": set(), "personas": set()})
        gd["minutos"] += j.minutos_presencia or 0
        gd["dias"].add(j.fecha)
        gd["personas"].add(j.persona_id)

        nom = f"{j.persona.apellido}, {j.persona.nombre}" if j.persona else "?"
        pd = por_persona.setdefault((g, nom), {"minutos": 0, "dias": 0})
        pd["minutos"] += j.minutos_presencia or 0
        pd["dias"] += 1 if (j.minutos_presencia or 0) else 0

    resumen = sorted(
        ({"grupo": g, "minutos": v["minutos"], "dias": len(v["dias"]),
          "personas": len(v["personas"])} for g, v in por_grupo.items()),
        key=lambda x: -x["minutos"])
    detalle = sorted(
        ({"grupo": g, "persona": n, **v} for (g, n), v in por_persona.items()),
        key=lambda x: (x["grupo"], -x["minutos"]))

    prev_m = (desde - timedelta(days=1)).replace(day=1)
    next_m = (hasta + timedelta(days=1))
    return templates.TemplateResponse(request, "asistencia/terceros.html", {
        "current_user": current_user,
        "anio": anio, "mes_num": m, "desde": desde, "hasta": hasta,
        "prev": prev_m.strftime("%Y-%m"), "next": next_m.strftime("%Y-%m"),
        "hay_next": next_m <= hoy,
        "resumen": resumen, "detalle": detalle,
        "total_min": sum(r["minutos"] for r in resumen),
        "total_personas": len({j.persona_id for j in jornadas}),
        "mt": min_a_texto,
    })


# ══════════════════════════════════════════════════════════════════════════════
# Dashboard mensual
# ══════════════════════════════════════════════════════════════════════════════

_reportes = require_perm("asistencia.reportes")


def _rango_mes(mes: str):
    hoy = date.today()
    try:
        anio, m = (int(x) for x in mes.split("-"))
        date(anio, m, 1)
    except (ValueError, TypeError):
        anio, m = hoy.year, hoy.month
    return anio, m, date(anio, m, 1), date(anio, m, monthrange(anio, m)[1])


def _datos_mes(db: Session, desde: date, hasta: date, planta: str = "", grupo: str = ""):
    q = (db.query(AsistenciaJornada)
         .options(joinedload(AsistenciaJornada.persona),
                  joinedload(AsistenciaJornada.motivo))
         .filter(AsistenciaJornada.fecha >= desde,
                 AsistenciaJornada.fecha <= hasta,
                 AsistenciaJornada.tipo_persona_snap != "tercero"))
    if planta:
        q = q.filter(AsistenciaJornada.planta_snap == planta)
    if grupo:
        q = q.filter(AsistenciaJornada.grupo_snap == grupo)
    return q.order_by(AsistenciaJornada.fecha).all()


@router.get("/mes", response_class=HTMLResponse)
async def mes(request: Request, db: Session = Depends(get_db),
              current_user=Depends(_reportes),
              mes: str = "", planta: str = "", grupo: str = ""):
    """Parcial del mes en curso y reporte final del mes cerrado.

    Es la misma pantalla: mientras el mes corre muestra hasta dónde llegó la
    carga, y cuando está completo es el reporte. Separarlas en dos daría dos
    números distintos para la misma pregunta.
    """
    cfg = get_config(db)
    hoy = date.today()
    anio, m, desde, hasta = _rango_mes(mes)
    jornadas = _datos_mes(db, desde, hasta, planta, grupo)

    feriados = {f.fecha for f in db.query(AsistenciaFeriado).filter(
        AsistenciaFeriado.fecha >= desde, AsistenciaFeriado.fecha <= hasta).all()}

    # ── Serie diaria ─────────────────────────────────────────────────────────
    por_dia = {}
    for j in jornadas:
        d = por_dia.setdefault(j.fecha, {"e50": 0, "e100": 0, "n": 0, "personas": set()})
        d["e50"] += j.minutos_extra_50 or 0
        d["e100"] += j.minutos_extra_100 or 0
        d["n"] += 1
        if (j.minutos_extra_detectada or 0) > 0:
            d["personas"].add(j.persona_id)

    dias, laborables, cargados = [], 0, 0
    for n in range(1, monthrange(anio, m)[1] + 1):
        f = date(anio, m, n)
        dat = por_dia.get(f)
        es_laborable = f.weekday() != 6 and f not in feriados
        if es_laborable and f <= hoy:
            laborables += 1
            if dat:
                cargados += 1
        dias.append({
            "fecha": f, "dia": n, "dow": f.weekday(),
            "e50": dat["e50"] if dat else 0,
            "e100": dat["e100"] if dat else 0,
            "total": (dat["e50"] + dat["e100"]) if dat else 0,
            "personas": len(dat["personas"]) if dat else 0,
            "cargado": bool(dat),
            "futuro": f > hoy,
            "feriado": f in feriados,
            "domingo": f.weekday() == 6,
        })

    # ── Ranking de personas ──────────────────────────────────────────────────
    porp = {}
    for j in jornadas:
        k = j.persona_id
        p = porp.setdefault(k, {
            "persona": j.persona, "e50": 0, "e100": 0, "aprobada": 0,
            "trabajado": 0, "defecto": 0, "dias_extra": 0, "fechas_extra": set(),
            "planta": j.planta_snap, "grupo": j.grupo_snap, "pendientes": 0,
        })
        p["e50"] += j.minutos_extra_50 or 0
        p["e100"] += j.minutos_extra_100 or 0
        p["trabajado"] += j.minutos_trabajados or 0
        p["defecto"] += j.minutos_defecto or 0
        if j.minutos_extra_aprobada is not None:
            p["aprobada"] += j.minutos_extra_aprobada
        if (j.minutos_extra_detectada or 0) > 0:
            p["dias_extra"] += 1
            p["fechas_extra"].add(j.fecha)
            if j.estado_extra == "pendiente":
                p["pendientes"] += 1

    def _racha(fechas):
        """Días consecutivos con extra más largos del mes."""
        if not fechas:
            return 0
        ord_f = sorted(fechas)
        mejor = actual = 1
        for a, b in zip(ord_f, ord_f[1:]):
            actual = actual + 1 if (b - a).days == 1 else 1
            mejor = max(mejor, actual)
        return mejor

    ranking = []
    for v in porp.values():
        total = v["e50"] + v["e100"]
        ranking.append({**v, "total": total, "racha": _racha(v["fechas_extra"])})
    ranking.sort(key=lambda x: -x["total"])

    # ── Cortes ───────────────────────────────────────────────────────────────
    def _agrupar(campo):
        acc = {}
        for j in jornadas:
            k = getattr(j, campo) or "—"
            a = acc.setdefault(k, {"clave": k, "total": 0, "personas": set()})
            a["total"] += (j.minutos_extra_50 or 0) + (j.minutos_extra_100 or 0)
            a["personas"].add(j.persona_id)
        out = [{"clave": v["clave"], "total": v["total"], "personas": len(v["personas"])}
               for v in acc.values()]
        return sorted(out, key=lambda x: -x["total"])

    e50 = sum(j.minutos_extra_50 or 0 for j in jornadas)
    e100 = sum(j.minutos_extra_100 or 0 for j in jornadas)
    total_extra = e50 + e100
    aprobada = sum(j.minutos_extra_aprobada or 0 for j in jornadas
                   if j.minutos_extra_aprobada is not None)

    estados = {}
    for j in jornadas:
        if (j.minutos_extra_detectada or 0) > 0:
            estados[j.estado_extra] = estados.get(j.estado_extra, 0) + 1

    # ── Umbrales de acumulación ──────────────────────────────────────────────
    umbrales = [(cfg.umbral_mes_3_min, 3), (cfg.umbral_mes_2_min, 2), (cfg.umbral_mes_1_min, 1)]
    mes_qs = f"{anio}-{m:02d}"
    prom = (sum(r["total"] for r in ranking) / len(ranking)) if ranking else 0
    alertas = []
    for r in ranking:
        motivos = []
        nivel = next((n for u, n in umbrales if r["total"] >= (u or 10 ** 9)), 0)
        if nivel:
            umbral = next(u for u, n in umbrales if n == nivel)
            motivos.append(f"acumuló {min_a_texto(r['total'])}, por encima del umbral de "
                           f"{min_a_texto(umbral)}")
        if r["racha"] >= (cfg.dias_racha_alerta or 3):
            motivos.append(f"{r['racha']} días seguidos con extra")
        if prom and r["total"] >= prom * 2 and r["total"] > 0:
            motivos.append(f"hizo {r['total'] / prom:.1f}× el promedio del mes")
        if r["pendientes"]:
            motivos.append(f"{r['pendientes']} día(s) sin resolver")
        if not motivos:
            continue
        alertas.append({**r, "nivel": nivel, "motivos": motivos,
                        "url": f"/asistencia/persona/{r['persona'].id}?mes={mes_qs}"})

    plantas_disp = sorted({j.planta_snap for j in _datos_mes(db, desde, hasta) if j.planta_snap})
    grupos_disp = sorted({j.grupo_snap for j in _datos_mes(db, desde, hasta) if j.grupo_snap})

    prev_m = (desde - timedelta(days=1)).replace(day=1)
    next_m = hasta + timedelta(days=1)
    return templates.TemplateResponse(request, "asistencia/mes.html", {
        "current_user": current_user,
        "anio": anio, "mes_num": m, "desde": desde, "hasta": hasta,
        "prev": prev_m.strftime("%Y-%m"), "next": next_m.strftime("%Y-%m"),
        "hay_next": next_m <= hoy, "hoy": hoy,
        "meses": _meses_con_datos(db),
        "en_curso": desde <= hoy <= hasta,
        "planta": planta, "grupo": grupo,
        "plantas_disp": plantas_disp, "grupos_disp": grupos_disp,
        "dias": dias, "ranking": ranking, "alertas": alertas,
        "por_grupo": _agrupar("grupo_snap"), "por_planta": _agrupar("planta_snap"),
        "estados": estados,
        "laborables": laborables, "cargados": cargados,
        "kpis": {
            "extra": total_extra, "e50": e50, "e100": e100, "aprobada": aprobada,
            "pendiente": total_extra - aprobada,
            "personas_extra": len([r for r in ranking if r["total"] > 0]),
            "personas": len(ranking),
            "dias_extra": len([d for d in dias if d["total"] > 0]),
            "trabajado": sum(j.minutos_trabajados or 0 for j in jornadas),
            "defecto": sum(j.minutos_defecto or 0 for j in jornadas),
            "max_dia": max([d["total"] for d in dias], default=0),
            # Portería y cualquier otro régimen rotativo: las horas se compensan
            # con francos, no generan extra. Se muestran igual — esconderlas
            # sin dejar rastro haría que los totales no cierren contra el papel.
            "rotativo": sum(j.minutos_presencia or 0 for j in jornadas
                            if not j.computa_extra_snap),
            "rotativo_personas": len({j.persona_id for j in jornadas
                                      if not j.computa_extra_snap}),
        },
        "mt": min_a_texto,
        "estado_extra_css": ESTADO_EXTRA_CSS,
    })


@router.get("/mes/export")
async def mes_export(db: Session = Depends(get_db), current_user=Depends(_reportes),
                     mes: str = "", planta: str = "", grupo: str = ""):
    """Excel del mes, una fila por persona y día. Respeta los filtros de pantalla."""
    import io as _io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    anio, m, desde, hasta = _rango_mes(mes)
    jornadas = _datos_mes(db, desde, hasta, planta, grupo)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{anio}-{m:02d}"
    cols = ["Fecha", "Día", "Persona", "Planta", "Grupo", "Ingreso", "Egreso",
            "Debe (h)", "Trabajado (h)", "Extra 50% (h)", "Extra 100% (h)",
            "Extra total (h)", "Aprobada (h)", "No cumplidas (h)",
            "Estado extra", "Registro", "Motivo", "Nota de la planilla"]
    ws.append(cols)
    cab = ws[1]
    for c in cab:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A8A")
        c.alignment = Alignment(horizontal="center")

    dias_sem = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
    h = lambda x: round((x or 0) / 60.0, 2)  # noqa: E731
    for j in sorted(jornadas, key=lambda x: (x.fecha, x.persona.apellido if x.persona else "")):
        ws.append([
            j.fecha, dias_sem[j.fecha.weekday()],
            f"{j.persona.apellido}, {j.persona.nombre}" if j.persona else "",
            j.planta_snap or "", j.grupo_snap or "",
            j.ingreso_real.strftime("%H:%M") if j.ingreso_real else "",
            j.egreso_real.strftime("%H:%M") if j.egreso_real else "",
            h(j.minutos_esperados), h(j.minutos_trabajados),
            h(j.minutos_extra_50), h(j.minutos_extra_100),
            h(j.minutos_extra_detectada),
            h(j.minutos_extra_aprobada) if j.minutos_extra_aprobada is not None else "",
            h(j.minutos_defecto),
            j.estado_extra or "", j.estado_registro or "",
            j.motivo.nombre if j.motivo else "", j.nota_origen or "",
        ])

    anchos = [11, 11, 26, 8, 16, 9, 9, 9, 12, 12, 13, 13, 12, 14, 13, 13, 20, 34]
    for i, w in enumerate(anchos, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    sufijo = ("_" + planta if planta else "") + ("_" + grupo.replace(" ", "") if grupo else "")
    nombre = f"asistencia_{anio}-{m:02d}{sufijo}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )


# ══════════════════════════════════════════════════════════════════════════════
# Buzón automático
# ══════════════════════════════════════════════════════════════════════════════

def _estado_buzon(db: Session) -> dict:
    """Todo lo que la pantalla necesita saber del buzón."""
    from app import asistencia_mail

    ultimas = (db.query(AsistenciaImportacion)
               .order_by(AsistenciaImportacion.created_at.desc())
               .limit(8).all())
    ultima_ok = (db.query(AsistenciaImportacion)
                 .filter(AsistenciaImportacion.origen == "buzon",
                         AsistenciaImportacion.estado.in_(["ok", "parcial"]))
                 .order_by(AsistenciaImportacion.created_at.desc()).first())

    dias_sin = None
    if ultima_ok and ultima_ok.created_at:
        dias_sin = (datetime.utcnow() - ultima_ok.created_at).days

    pendientes = []
    for r in ultimas:
        if r.dias_pendientes and r.detalle:
            try:
                pendientes += json.loads(r.detalle).get("pendientes", [])
            except (ValueError, TypeError):
                pass

    return {
        "configurado": asistencia_mail.configurado(),
        "cfg": asistencia_mail.config(),
        "ultimas": ultimas,
        "ultima_ok": ultima_ok,
        "dias_sin_planilla": dias_sin,
        "pendientes_revision": pendientes[:10],
    }


@router.post("/importar/buzon")
async def importar_buzon(db: Session = Depends(get_db), current_user=Depends(_cargar)):
    """Traer ahora: revisa el buzón y aplica lo que sea seguro aplicar."""
    from app.asistencia_mail import procesar_buzon

    regs = procesar_buzon(db, usuario_id=current_user.id, origen="buzon")
    if not regs:
        return RedirectResponse(
            "/asistencia/importar?ok=No+hay+planillas+nuevas+en+el+buzón", 303)

    err = next((r for r in regs if r.estado == "error"), None)
    if err is not None:
        return RedirectResponse(
            f"/asistencia/importar?err={(err.error or 'Error')[:160]}", 303)

    dias = sum(r.dias_aplicados for r in regs)
    filas = sum(r.filas_aplicadas for r in regs)
    pend = sum(r.dias_pendientes for r in regs)
    msg = f"{dias}+días+importados+·+{filas}+filas"
    if pend:
        msg += f"+·+{pend}+días+para+revisar"
    return RedirectResponse(f"/asistencia?ok={msg}", 303)


@router.post("/importar/buzon/probar")
async def importar_buzon_probar(db: Session = Depends(get_db),
                                current_user=Depends(_cargar)):
    from app.asistencia_mail import probar_conexion

    r = probar_conexion()
    if r.get("ok"):
        return RedirectResponse(
            f"/asistencia/importar?ok=Conexión+OK+·+{r['user']}+·+"
            f"{r['mensajes']}+mensajes+en+{r['carpeta']}", 303)
    extra = ""
    if r.get("carpetas"):
        extra = "+·+Carpetas:+" + ",+".join(r["carpetas"][:6])
    return RedirectResponse(
        f"/asistencia/importar?err={(r.get('error') or 'Error')[:160]}{extra}", 303)


# ══════════════════════════════════════════════════════════════════════════════
# Ficha individual
# ══════════════════════════════════════════════════════════════════════════════

_FRANJAS = [(0, 6, "Noche 00-06"), (6, 12, "Mañana 06-12"),
            (12, 18, "Tarde 12-18"), (18, 24, "Noche 18-00")]


def _meses_con_datos(db: Session, persona_id: int = None) -> list:
    """Meses que tienen jornadas cargadas, del más nuevo al más viejo.

    Alimenta el selector: no tiene sentido ofrecer meses vacíos, y a medida que
    se carguen octubre, noviembre… aparecen solos.
    """
    q = db.query(AsistenciaJornada.fecha)
    if persona_id:
        q = q.filter(AsistenciaJornada.persona_id == persona_id)
    vistos = {(f.year, f.month) for (f,) in q.distinct().all()}
    return sorted(vistos, reverse=True)


@router.get("/persona/{persona_id}", response_class=HTMLResponse)
async def persona_ficha(persona_id: int, request: Request,
                        db: Session = Depends(get_db),
                        current_user=Depends(_reportes), mes: str = "",
                        foco: str = ""):
    p = db.query(AsistenciaPersona).get(persona_id)
    if p is None:
        return RedirectResponse("/asistencia/mes?err=Persona+inexistente", 303)

    hoy = date.today()
    anio, m, desde, hasta = _rango_mes(mes)

    jornadas = (db.query(AsistenciaJornada)
                .filter(AsistenciaJornada.persona_id == p.id,
                        AsistenciaJornada.fecha >= desde,
                        AsistenciaJornada.fecha <= hasta)
                .order_by(AsistenciaJornada.fecha).all())
    bloques = {}
    for b in (db.query(AsistenciaBloque)
              .filter(AsistenciaBloque.persona_id == p.id,
                      AsistenciaBloque.fecha_operativa >= desde,
                      AsistenciaBloque.fecha_operativa <= hasta)
              .order_by(AsistenciaBloque.desde).all()):
        bloques.setdefault(b.fecha_operativa, []).append(b)
    for j in jornadas:
        j._bloques = bloques.get(j.fecha, [])

    trabajados = [j for j in jornadas if (j.minutos_presencia or 0) > 0]
    con_extra = [j for j in jornadas if (j.minutos_extra_detectada or 0) > 0]
    e50 = sum(j.minutos_extra_50 or 0 for j in jornadas)
    e100 = sum(j.minutos_extra_100 or 0 for j in jornadas)
    trabajado = sum(j.minutos_trabajados or 0 for j in jornadas)
    esperado = sum(j.minutos_esperados or 0 for j in jornadas if j.laborable_esp)

    # ── Patrón horario ───────────────────────────────────────────────────────
    # Si alguien entra siempre a la misma hora, su horario real es ese. La
    # dispersión es lo que delata que el horario de la nómina no es el que rige.
    entradas, salidas, franjas = {}, {}, {n: 0 for _, _, n in _FRANJAS}
    for j in jornadas:
        for b in j._bloques:
            h = b.desde.hour
            entradas[h] = entradas.get(h, 0) + 1
            for ini, fin, nom in _FRANJAS:
                if ini <= h < fin:
                    franjas[nom] += 1
                    break
            if b.hasta:
                salidas[b.hasta.hour] = salidas.get(b.hasta.hour, 0) + 1

    def _moda(d):
        return max(d.items(), key=lambda kv: kv[1]) if d else (None, 0)
    ent_moda, ent_n = _moda(entradas)
    sal_moda, sal_n = _moda(salidas)

    # ── Día de la semana ─────────────────────────────────────────────────────
    por_dow = {i: {"extra": 0, "dias": 0} for i in range(7)}
    for j in jornadas:
        d = por_dow[j.fecha.weekday()]
        d["extra"] += j.minutos_extra_detectada or 0
        d["dias"] += 1 if (j.minutos_presencia or 0) else 0

    # ── Racha de días consecutivos con extra ─────────────────────────────────
    fechas_extra = sorted(j.fecha for j in con_extra)
    racha = actual = 1 if fechas_extra else 0
    for a, b in zip(fechas_extra, fechas_extra[1:]):
        actual = actual + 1 if (b - a).days == 1 else 1
        racha = max(racha, actual)

    # ── Comparación con sus pares ────────────────────────────────────────────
    # Es lo que convierte un número en un juicio: 8 h de extra no dicen nada
    # hasta saber que sus compañeros hicieron 2.
    pares_q = (db.query(AsistenciaJornada)
               .filter(AsistenciaJornada.fecha >= desde,
                       AsistenciaJornada.fecha <= hasta,
                       AsistenciaJornada.persona_id != p.id,
                       AsistenciaJornada.tipo_persona_snap == (p.tipo or "mtr")))
    if p.grupo:
        pares_q = pares_q.filter(AsistenciaJornada.grupo_snap == p.grupo)
    por_par = {}
    for j in pares_q.all():
        d = por_par.setdefault(j.persona_id, {"extra": 0, "trab": 0})
        d["extra"] += j.minutos_extra_detectada or 0
        d["trab"] += j.minutos_trabajados or 0
    n_pares = len(por_par)
    prom_extra = (sum(v["extra"] for v in por_par.values()) / n_pares) if n_pares else 0
    prom_trab = (sum(v["trab"] for v in por_par.values()) / n_pares) if n_pares else 0
    extra_total = e50 + e100
    peores = sorted((v["extra"] for v in por_par.values()), reverse=True)
    puesto = sum(1 for x in peores if x > extra_total) + 1

    # ── Evolución de los últimos 6 meses ─────────────────────────────────────
    serie = []
    y, mm = anio, m
    for _ in range(6):
        d1 = date(y, mm, 1)
        d2 = date(y, mm, monthrange(y, mm)[1])
        tot = (db.query(func.coalesce(func.sum(AsistenciaJornada.minutos_extra_detectada), 0))
               .filter(AsistenciaJornada.persona_id == p.id,
                       AsistenciaJornada.fecha >= d1,
                       AsistenciaJornada.fecha <= d2).scalar() or 0)
        hay = (db.query(func.count(AsistenciaJornada.id))
               .filter(AsistenciaJornada.persona_id == p.id,
                       AsistenciaJornada.fecha >= d1,
                       AsistenciaJornada.fecha <= d2).scalar() or 0)
        serie.append({"anio": y, "mes": mm, "extra": tot, "hay": bool(hay)})
        mm -= 1
        if mm == 0:
            y, mm = y - 1, 12
    serie.reverse()

    # ── Alertas de esta persona ──────────────────────────────────────────────
    cfg = get_config(db)
    alertas = []
    for umbral, nivel in [(cfg.umbral_mes_3_min, 3), (cfg.umbral_mes_2_min, 2),
                          (cfg.umbral_mes_1_min, 1)]:
        if umbral and extra_total >= umbral:
            alertas.append(("critica" if nivel == 3 else "media",
                            f"Acumuló {min_a_texto(extra_total)} — supera el umbral de "
                            f"{min_a_texto(umbral)}"))
            break
    if racha >= (cfg.dias_racha_alerta or 3):
        alertas.append(("media", f"{racha} días seguidos con horas extra"))
    if trabajados and len(con_extra) * 100 / len(trabajados) >= (cfg.pct_recurrencia_alerta or 60):
        alertas.append(("media",
                        f"Hizo extra en {len(con_extra)} de {len(trabajados)} días trabajados "
                        f"({len(con_extra) * 100 // len(trabajados)} %) — su horario real "
                        f"puede no ser el de la nómina"))
    incompletas = [j for j in jornadas if j.estado_registro in ("sin_egreso", "sin_ingreso",
                                                                "inconsistente")]
    if incompletas:
        alertas.append(("gris", f"{len(incompletas)} día(s) con el registro incompleto — "
                                f"esas horas no computan"))
    largas = [j for j in jornadas
              if (j.minutos_presencia or 0) > (cfg.umbral_jornada_larga_min or 720)]
    if largas:
        alertas.append(("critica", f"{len(largas)} jornada(s) de más de "
                                   f"{min_a_texto(cfg.umbral_jornada_larga_min)}"))
    if n_pares and prom_extra and extra_total >= prom_extra * 2:
        alertas.append(("media", f"Hizo {extra_total / prom_extra:.1f}× la extra promedio "
                                 f"de su grupo"))

    prev_m = (desde - timedelta(days=1)).replace(day=1)
    next_m = hasta + timedelta(days=1)
    return templates.TemplateResponse(request, "asistencia/persona.html", {
        # La clave es "persona" y no "p": base.html define `p` como la ruta
        # actual dentro del bloque content y pisaría a la persona.
        "current_user": current_user, "persona": p,
        "anio": anio, "mes_num": m, "desde": desde, "hasta": hasta,
        "prev": prev_m.strftime("%Y-%m"), "next": next_m.strftime("%Y-%m"),
        "hay_next": next_m <= hoy,
        "meses": _meses_con_datos(db),
        "jornadas": jornadas, "alertas": alertas,
        "foco": _parse_fecha(foco, None) if foco else None,
        "motivos": db.query(AsistenciaMotivo)
                     .filter(AsistenciaMotivo.activo == True)  # noqa: E712
                     .order_by(AsistenciaMotivo.orden).all(),
        "puede_aprobar": can(current_user, "asistencia.aprobar_extra"),
        "min_a_hhmm": min_a_hhmm,
        "entradas": entradas, "salidas": salidas, "franjas": franjas,
        "ent_moda": ent_moda, "ent_n": ent_n, "sal_moda": sal_moda, "sal_n": sal_n,
        "por_dow": por_dow, "serie": serie,
        "comparacion": {
            "n_pares": n_pares, "prom_extra": prom_extra, "prom_trab": prom_trab,
            "puesto": puesto, "total": n_pares + 1,
            "ratio": (extra_total / prom_extra) if prom_extra else None,
        },
        "kpis": {
            "extra": extra_total, "e50": e50, "e100": e100,
            "aprobada": sum(j.minutos_extra_aprobada or 0 for j in jornadas
                            if j.minutos_extra_aprobada is not None),
            "trabajado": trabajado, "esperado": esperado,
            "cumplimiento": round(trabajado * 100 / esperado) if esperado else None,
            "defecto": sum(j.minutos_defecto or 0 for j in jornadas),
            "dias_trabajados": len(trabajados),
            "dias_extra": len(con_extra),
            "ausentes": len([j for j in jornadas if j.estado_registro == "ausente"]),
            "incompletas": len(incompletas),
            "racha": racha,
            "pct_extra": round(len(con_extra) * 100 / len(trabajados)) if trabajados else 0,
        },
        "mt": min_a_texto,
        "estado_extra_css": ESTADO_EXTRA_CSS,
        "estado_registro_css": ESTADO_REGISTRO_CSS,
    })


# ══════════════════════════════════════════════════════════════════════════════
# Resolución de horas extra
# ══════════════════════════════════════════════════════════════════════════════

_aprobar = require_perm("asistencia.aprobar_extra")


def _ctx_jornada(db, j):
    return {
        "j": j,
        "motivos": db.query(AsistenciaMotivo)
                     .filter(AsistenciaMotivo.activo == True)  # noqa: E712
                     .order_by(AsistenciaMotivo.orden).all(),
        "puede_aprobar": True,
        "mt": min_a_texto,
        "estado_extra_css": ESTADO_EXTRA_CSS,
        "estado_registro_css": ESTADO_REGISTRO_CSS,
        "min_a_hhmm": min_a_hhmm,
    }


def _resolver(db, j, accion, motivo_id, detalle, minutos, comentario, user):
    """Aplica una decisión sobre la extra de una jornada. Devuelve (ok, error).

    El dato de asistencia NO se toca: se decide sobre la extra detectada, que
    sigue siendo la que sale del cálculo. Si lo que está mal es el horario, se
    corrige el horario y el sistema recalcula.
    """
    cfg = get_config(db)
    if (j.minutos_extra_detectada or 0) <= 0:
        return False, "Esa jornada no tiene horas extra para resolver."

    if accion == "rechazar":
        nuevos = 0
    else:
        nuevos = (hhmm_a_min(minutos) if (minutos or "").strip()
                  else j.minutos_extra_detectada)
        if nuevos is None:
            return False, "La cantidad de horas no es válida. Usá el formato HH:MM."

    mid = int(motivo_id) if (motivo_id or "").strip().isdigit() else None
    if cfg.motivo_obligatorio_aprobar and accion == "aprobar" and not mid:
        return False, "Falta el motivo: sin él, el análisis de causas no sirve."

    # Cambiar lo aprobado respecto de lo detectado exige explicación. Es el
    # único lugar donde el sistema se pone pesado, y es donde corresponde.
    if accion == "aprobar" and nuevos != j.minutos_extra_detectada \
            and not (comentario or "").strip():
        return False, (f"Estás aprobando {min_a_texto(nuevos)} sobre "
                       f"{min_a_texto(j.minutos_extra_detectada)} detectadas. "
                       f"Escribí por qué.")

    motivo = db.query(AsistenciaMotivo).get(mid) if mid else None
    if motivo is not None and motivo.requiere_detalle and not (detalle or "").strip():
        return False, f"El motivo «{motivo.nombre}» pide un detalle."

    for campo, viejo, nuevo in [
        ("estado_extra", j.estado_extra, "aprobada" if accion == "aprobar" else "rechazada"),
        ("minutos_extra_aprobada", j.minutos_extra_aprobada, nuevos),
        ("motivo_id", j.motivo_id, mid),
        ("motivo_detalle", j.motivo_detalle, (detalle or "").strip() or None),
    ]:
        if viejo != nuevo:
            setattr(j, campo, nuevo)
            db.add(AsistenciaAuditLog(
                entidad="jornada", entidad_id=j.id, accion=accion, campo=campo,
                valor_anterior=str(viejo) if viejo is not None else None,
                valor_nuevo=str(nuevo) if nuevo is not None else None,
                motivo=(comentario or "").strip() or None, user_id=user.id))

    j.comentario_aprobacion = (comentario or "").strip() or None
    j.aprobado_por_id = user.id
    j.aprobado_at = datetime.utcnow()
    j.revisar = False        # la decisión vuelve a estar al día
    return True, None


@router.post("/jornada/{jornada_id}/resolver", response_class=HTMLResponse)
async def resolver_jornada(jornada_id: int, request: Request,
                           db: Session = Depends(get_db), current_user=Depends(_aprobar),
                           accion: str = Form("aprobar"), motivo_id: str = Form(""),
                           detalle: str = Form(""), minutos: str = Form(""),
                           comentario: str = Form("")):
    j = db.query(AsistenciaJornada).get(jornada_id)
    if j is None:
        return HTMLResponse("", status_code=404)

    if accion == "reabrir":
        db.add(AsistenciaAuditLog(
            entidad="jornada", entidad_id=j.id, accion="reabrir", campo="estado_extra",
            valor_anterior=j.estado_extra, valor_nuevo="pendiente", user_id=current_user.id))
        j.estado_extra = "pendiente"
        j.minutos_extra_aprobada = None
        j.aprobado_por_id, j.aprobado_at = None, None
        db.commit()
        db.refresh(j)
        ctx = _ctx_jornada(db, j)
        ctx["current_user"] = current_user
        return templates.TemplateResponse(request, "asistencia/_jornada_fila.html", ctx)

    ok, err = _resolver(db, j, accion, motivo_id, detalle, minutos, comentario, current_user)
    if not ok:
        db.rollback()
        db.refresh(j)
        ctx = _ctx_jornada(db, j)
        ctx.update({"current_user": current_user, "error": err})
        return templates.TemplateResponse(request, "asistencia/_jornada_fila.html", ctx)

    db.commit()
    db.refresh(j)
    ctx = _ctx_jornada(db, j)
    ctx["current_user"] = current_user
    return templates.TemplateResponse(request, "asistencia/_jornada_fila.html", ctx)


@router.post("/resolver-lote")
async def resolver_lote(db: Session = Depends(get_db), current_user=Depends(_aprobar),
                        jornadas: list = Form([]), accion: str = Form("aprobar"),
                        motivo_id: str = Form(""), detalle: str = Form(""),
                        comentario: str = Form(""), volver: str = Form("/asistencia")):
    """Resuelve varias jornadas de una. Sin esto nadie aprueba 40 extras al mes
    de a una, y el estado 'pendiente' se vuelve decorativo."""
    ids = [int(x) for x in jornadas if str(x).isdigit()]
    if not ids:
        return RedirectResponse(f"{volver}{'&' if '?' in volver else '?'}"
                                f"err=No+elegiste+ninguna+jornada", 303)

    hechas, fallidas, primer_error = 0, 0, None
    for j in db.query(AsistenciaJornada).filter(AsistenciaJornada.id.in_(ids)).all():
        ok, err = _resolver(db, j, accion, motivo_id, detalle, "", comentario, current_user)
        if ok:
            hechas += 1
        else:
            fallidas += 1
            primer_error = primer_error or err
    db.commit()

    sep = "&" if "?" in volver else "?"
    verbo = "aprobadas" if accion == "aprobar" else "rechazadas"
    if hechas and not fallidas:
        return RedirectResponse(f"{volver}{sep}ok={hechas}+jornadas+{verbo}", 303)
    if hechas:
        return RedirectResponse(
            f"{volver}{sep}ok={hechas}+{verbo}+·+{fallidas}+sin+resolver:+{primer_error[:80]}", 303)
    return RedirectResponse(f"{volver}{sep}err={(primer_error or 'No se pudo')[:140]}", 303)

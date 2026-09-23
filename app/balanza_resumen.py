"""
Lector del resumen de descarga que manda balanza cuando termina un buque.

QUÉ ES ESE ARCHIVO
    Un Excel con TODOS los pesajes del operativo, uno por renglón: número de
    ticket, patente, hora de entrada y de salida, tara, bruto, neto y el peso
    declarado en origen. Viene en dos hojas con los mismos tickets agrupados de
    distinta forma — la primera por turno, la segunda por transportista — más
    un encabezado con el buque, el cliente, el producto y las fechas.

    Es el registro más fino que existe de un operativo: de acá sale el ritmo
    real, quién movió cuánto, cuánto tarda un camión adentro de la planta y
    cuánto se despega nuestra balanza de la del origen.

CÓMO SE RECONOCE
    Balanza manda muchos Excel a la misma casilla: horas del personal,
    quincenas de CNA y TYS, tolling, stock. El asunto no sirve para
    distinguirlos — cambia todo el tiempo y a veces ni nombra el buque.

    Lo que sí distingue es el archivo por dentro: un resumen de buque tiene
    "Operativo: NN) NOMBRE" en el encabezado y una columna "Pat. Cam.". Las
    quincenas tienen la columna pero no el encabezado; las horas no tienen
    ninguna de las dos. El "NN)" importa: los operativos de buque van
    numerados, y sin ese número entran cosas como
    "Operativo: SEPOR S.A. (SAUSOR)", que no es un buque.

LO QUE NO HACE
    No decide si el operativo terminó ni lo cruza con nada. Devuelve lo que
    leyó. Un archivo que no es un resumen de buque devuelve None, sin error.
"""
from __future__ import annotations

import io
import re
import unicodedata
from datetime import date, datetime, time, timedelta

from openpyxl import load_workbook

# Encabezado que sólo tienen los resúmenes de buque.
_OPERATIVO = re.compile(r"Operativo:\s*(\d+)\s*\)\s*(.+)", re.I)
_ETIQUETA = re.compile(r"^\s*([^:]{3,30}):\s*(.*)$")
_FECHA = re.compile(r"(\d{1,2})[-/](\d{1,2})[-/](\d{2,4})")

# Cómo se llaman las columnas en la fila de encabezado de cada bloque.
_COLS = {"Nro.": "nro", "Fecha Ent.": "f_ent", "H. Ent.": "h_ent",
         "Fecha Sal.": "f_sal", "H. Sal.": "h_sal", "Pat. Cam.": "patente",
         "Tara": "tara", "Bruto": "bruto", "Neto": "neto",
         "Peso Orig.": "origen", "Diferencia": "diferencia",
         "Cliente": "cliente", "Producto": "producto"}


def _txt(v) -> str:
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def _sin_acentos(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "")
    return "".join(c for c in s if not unicodedata.combining(c)).upper()


def _entero(v) -> int | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v)
    s = re.sub(r"[^\d-]", "", str(v or ""))
    return int(s) if s and s.strip("-") else None


def _fecha(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = _FECHA.search(str(v or ""))
    if not m:
        return None
    d, mes, a = int(m.group(1)), int(m.group(2)), int(m.group(3))
    try:
        return date(a if a > 99 else 2000 + a, mes, d)
    except ValueError:
        return None


def _hora(v) -> time | None:
    if isinstance(v, time):
        return v
    if isinstance(v, datetime):
        return v.time()
    m = re.search(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", str(v or ""))
    if not m:
        return None
    try:
        return time(int(m.group(1)), int(m.group(2)), int(m.group(3) or 0))
    except ValueError:
        return None


def _cabecera(ws) -> dict:
    """Los datos sueltos del encabezado: operativo, fechas, cliente, producto."""
    out = {}
    for fila in ws.iter_rows(min_row=1, max_row=14, values_only=True):
        for celda in fila:
            s = _txt(celda)
            if not s:
                continue
            m = _OPERATIVO.match(s)
            if m:
                out["operativo_nro"] = int(m.group(1))
                out["buque"] = m.group(2).strip()
                continue
            e = _ETIQUETA.match(s)
            if not e:
                continue
            clave, valor = _sin_acentos(e.group(1)), e.group(2).strip()
            if clave == "FECHA INICIO":
                out["inicio"] = _fecha(valor)
            elif clave == "FECHA FINALIZACION":
                out["fin"] = _fecha(valor)
            elif clave == "CLIENTE":
                out.setdefault("cliente", valor or None)
            elif clave == "PRODUCTO":
                out.setdefault("producto", valor or None)
    return out


def _mapa_de_columnas(fila) -> dict | None:
    """Posición de cada columna, leída de la fila de encabezado del bloque."""
    pos = {}
    for i, celda in enumerate(fila):
        nombre = _COLS.get(_txt(celda))
        if nombre:
            pos[nombre] = i
    return pos if {"nro", "patente", "neto"} <= set(pos) else None


def _tickets(ws, etiqueta_grupo: str) -> list[dict]:
    """Los pesajes de una hoja, con el grupo al que pertenece cada uno.

    `etiqueta_grupo` es "TURNO" o "TRANSPORTE": es lo que cambia entre las dos
    hojas, que por lo demás son idénticas.
    """
    pos, grupo, out = None, None, []
    for fila in ws.iter_rows(values_only=True):
        primera = _txt(fila[0]) if fila else ""

        nuevo = _mapa_de_columnas(fila)
        if nuevo:
            pos = nuevo
            continue

        e = _ETIQUETA.match(primera)
        if e and _sin_acentos(e.group(1)) == etiqueta_grupo:
            grupo = e.group(2).strip() or None
            continue
        # "Viajes: 5 | Neto: ..." cierra el grupo; "Totales Generales" también.
        if primera.upper().startswith(("VIAJES:", "TOTALES")):
            continue

        if pos is None:
            continue
        nro = _entero(fila[pos["nro"]]) if len(fila) > pos["nro"] else None
        patente = _txt(fila[pos["patente"]]) if len(fila) > pos["patente"] else ""
        if nro is None or not patente:
            continue

        def v(campo):
            i = pos.get(campo)
            return fila[i] if i is not None and len(fila) > i else None

        out.append({
            "nro": nro, "patente": patente.upper(), "grupo": grupo,
            "f_ent": _fecha(v("f_ent")), "h_ent": _hora(v("h_ent")),
            "f_sal": _fecha(v("f_sal")), "h_sal": _hora(v("h_sal")),
            "tara": _entero(v("tara")), "bruto": _entero(v("bruto")),
            "neto": _entero(v("neto")), "origen": _entero(v("origen")),
            "cliente": _txt(v("cliente")) or None,
            "producto": _txt(v("producto")) or None,
        })
    return out


def _sin_repetidos(tickets: list[dict]) -> list[dict]:
    """Un pesaje por número, quedándose con el que dice a qué grupo pertenece.

    El mismo ticket puede aparecer dos veces: una bajo su turno y otra bajo la
    sección "Todos". Gana el que nombra el turno real, que es el que sirve.
    """
    mejores = {}
    for t in tickets:
        generico = not t["grupo"] or _sin_acentos(t["grupo"]) == "TODOS"
        previo = mejores.get(t["nro"])
        if previo is None:
            mejores[t["nro"]] = t
        elif generico:
            continue
        elif not previo["grupo"] or _sin_acentos(previo["grupo"]) == "TODOS":
            mejores[t["nro"]] = t
    return list(mejores.values())


def _mas_frecuente(valores) -> str | None:
    vistos = {}
    for v in valores:
        if v:
            vistos[v] = vistos.get(v, 0) + 1
    return max(vistos, key=vistos.get) if vistos else None


def limpiar_buque(crudo: str, cliente: str | None = None) -> str:
    """El nombre del buque, sin la fecha ni el cliente que le pega balanza.

    El encabezado del operativo mezcla cosas: "MV TAI HONOR (28-06-26)",
    "OCEAN INNOVATION 21.09.26", "MV SOFIA MOVIPORT (MAP) 14.07.2026". Sin
    limpiarlo, el mismo buque en dos operativos no se reconoce como el mismo y
    tampoco se puede cruzar con Próximos Arribos.
    """
    s = re.sub(r"\([^)]*\)", " ", crudo or "")              # lo que va entre paréntesis
    s = re.sub(r"\b\d{1,2}[-./]\d{1,2}[-./]\d{2,4}\b", " ", s)  # fechas pegadas
    s = re.sub(r"[^A-Za-z0-9ÁÉÍÓÚÜÑáéíóúüñ/. ]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip(" .-")
    # "MV SOFIA MOVIPORT" — el cliente quedó pegado al nombre.
    if cliente:
        primera = _sin_acentos(cliente).split()[0] if cliente.split() else ""
        if primera and len(primera) > 3 and _sin_acentos(s).endswith(" " + primera):
            s = s[: -(len(primera) + 1)].strip()
    return s


def _momento(f: date | None, h: time | None) -> datetime | None:
    return datetime.combine(f, h) if f and h else None


def parsear_resumen(datos: bytes, archivo: str = "") -> dict | None:
    """Lee el Excel y devuelve el operativo, o None si no es un resumen de buque.

    {"buque", "operativo_nro", "cliente", "producto", "inicio", "fin",
     "tickets": [{"nro", "patente", "turno", "transporte", "entrada",
                  "salida", "tara", "bruto", "neto", "origen"}],
     "avisos": [str]}
    """
    try:
        wb = load_workbook(io.BytesIO(datos), data_only=True)
    except Exception:
        return None

    try:
        cab = _cabecera(wb.worksheets[0])
        if not cab.get("buque"):
            return None                      # no es un resumen de buque

        por_turno = _tickets(wb.worksheets[0], "TURNO")
        if not por_turno:
            return None

        # El encabezado a veces dice "Producto: Todos" aunque el operativo
        # descargó uno solo: cada pesaje trae el suyo y ése es el bueno.
        if not cab.get("producto") or _sin_acentos(cab["producto"]) == "TODOS":
            cab["producto"] = _mas_frecuente(t["producto"] for t in por_turno)
        if not cab.get("cliente") or _sin_acentos(cab["cliente"]) == "TODOS":
            cab["cliente"] = _mas_frecuente(t["cliente"] for t in por_turno)
        cab["buque_crudo"] = cab["buque"]
        cab["buque"] = limpiar_buque(cab["buque"], cab.get("cliente"))
        # La segunda hoja trae los mismos tickets agrupados por transportista.
        por_transporte = []
        for ws in wb.worksheets[1:]:
            por_transporte += _tickets(ws, "TRANSPORTE")
        transporte_de = {t["nro"]: t["grupo"] for t in por_transporte if t["grupo"]}

        # Algunos archivos traen el detalle por turno y además una sección
        # "Turno: Todos" que repite todo. Sin deduplicar, ese operativo cuenta
        # el doble de viajes y el doble de toneladas.
        por_turno = _sin_repetidos(por_turno)
        por_transporte = _sin_repetidos(por_transporte)

        avisos = []
        sueltos = {t["nro"] for t in por_turno} ^ {t["nro"] for t in por_transporte}
        if por_transporte and sueltos:
            avisos.append(f"{len(sueltos)} pesaje(s) figuran en una hoja y no en la otra.")

        tickets = []
        for t in por_turno:
            entrada = _momento(t["f_ent"], t["h_ent"])
            salida = _momento(t["f_sal"], t["h_sal"])
            # El Excel a veces trae mal la fecha de salida de un camión que
            # entró antes de medianoche y salió después: queda un día atrás.
            if entrada and salida and salida < entrada:
                if salida + timedelta(days=1) - entrada < timedelta(hours=12):
                    salida += timedelta(days=1)
                else:
                    salida = None
            tickets.append({
                "nro": t["nro"], "patente": t["patente"],
                "turno": t["grupo"], "transporte": transporte_de.get(t["nro"]),
                "entrada": entrada, "salida": salida,
                "tara": t["tara"], "bruto": t["bruto"],
                "neto": t["neto"], "origen": t["origen"],
            })

        sin_transporte = sum(1 for t in tickets if not t["transporte"])
        if sin_transporte:
            avisos.append(f"{sin_transporte} pesaje(s) sin transportista identificado.")

        return {**cab, "archivo": archivo, "tickets": tickets, "avisos": avisos}
    finally:
        wb.close()
